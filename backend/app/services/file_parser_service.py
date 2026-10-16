import csv
import io
import json
from dataclasses import dataclass
from pathlib import Path


@dataclass
class ParsedFile:
    content: str
    mime_type: str
    extension: str


class FileParserService:
    SUPPORTED_EXTENSIONS = {
        ".txt",
        ".md",
        ".markdown",
        ".csv",
        ".json",
        ".log",
        ".pdf",
        ".docx",
        ".xlsx",
    }

    def parse(self, *, file_name: str, mime_type: str | None, raw_bytes: bytes) -> ParsedFile:
        extension = Path(file_name).suffix.lower()
        detected_mime = mime_type or self._guess_mime(extension)
        if extension not in self.SUPPORTED_EXTENSIONS:
            raise ValueError(f"Unsupported file type: {extension or 'unknown'}")

        if extension in {".txt", ".md", ".markdown", ".log"}:
            content = self._decode_text(raw_bytes)
        elif extension == ".json":
            content = self._format_json(raw_bytes)
        elif extension == ".csv":
            content = self._parse_csv(raw_bytes)
        elif extension == ".pdf":
            content = self._parse_pdf(raw_bytes)
        elif extension == ".docx":
            content = self._parse_docx(raw_bytes)
        elif extension == ".xlsx":
            content = self._parse_xlsx(raw_bytes)
        else:
            raise ValueError(f"Unsupported file type: {extension}")

        normalized = content.strip()
        if not normalized:
            raise ValueError("The uploaded file did not contain any readable text.")
        return ParsedFile(content=normalized, mime_type=detected_mime, extension=extension)

    @staticmethod
    def _decode_text(raw_bytes: bytes) -> str:
        encodings = ("utf-8", "utf-8-sig", "gb18030", "gbk")
        for encoding in encodings:
            try:
                return raw_bytes.decode(encoding)
            except UnicodeDecodeError:
                continue
        return raw_bytes.decode("latin-1")

    def _format_json(self, raw_bytes: bytes) -> str:
        payload = json.loads(self._decode_text(raw_bytes))
        return json.dumps(payload, ensure_ascii=False, indent=2)

    def _parse_csv(self, raw_bytes: bytes) -> str:
        text = self._decode_text(raw_bytes)
        reader = csv.reader(io.StringIO(text))
        rows = []
        for row in reader:
            rows.append(" | ".join(cell.strip() for cell in row))
        return "\n".join(rows)

    def _parse_pdf(self, raw_bytes: bytes) -> str:
        try:
            from pypdf import PdfReader
        except ImportError as exc:
            raise ValueError("PDF parsing requires pypdf. Please install dependencies first.") from exc

        reader = PdfReader(io.BytesIO(raw_bytes))
        texts = [(page.extract_text() or "").strip() for page in reader.pages]
        return "\n\n".join(item for item in texts if item)

    def _parse_docx(self, raw_bytes: bytes) -> str:
        try:
            from docx import Document
        except ImportError as exc:
            raise ValueError("DOCX parsing requires python-docx. Please install dependencies first.") from exc

        document = Document(io.BytesIO(raw_bytes))
        parts = [paragraph.text.strip() for paragraph in document.paragraphs if paragraph.text.strip()]
        return "\n".join(parts)

    def _parse_xlsx(self, raw_bytes: bytes) -> str:
        try:
            from openpyxl import load_workbook
        except ImportError as exc:
            raise ValueError("XLSX parsing requires openpyxl. Please install dependencies first.") from exc

        workbook = load_workbook(io.BytesIO(raw_bytes), data_only=True)
        lines: list[str] = []
        for sheet in workbook.worksheets:
            lines.append(f"# Sheet: {sheet.title}")
            for row in sheet.iter_rows(values_only=True):
                values = [str(cell).strip() for cell in row if cell is not None and str(cell).strip()]
                if values:
                    lines.append(" | ".join(values))
        return "\n".join(lines)

    @staticmethod
    def _guess_mime(extension: str) -> str:
        return {
            ".txt": "text/plain",
            ".md": "text/markdown",
            ".markdown": "text/markdown",
            ".csv": "text/csv",
            ".json": "application/json",
            ".log": "text/plain",
            ".pdf": "application/pdf",
            ".docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            ".xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        }.get(extension, "application/octet-stream")


file_parser_service = FileParserService()
